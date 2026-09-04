"""
================================================================================
Project     : OnlyNativeIQ
Organization: OnlyNative
Module      : Document Loader
File        : document_loader.py
Author      : Abhay Kumar Pandey
Created On  : 02-Sep-2026
Updated On  : 04-Sep-2026
Description :
    Loads authoritative knowledge documents from the OnlyNativeIQ knowledge
    base. Only PDF, DOCX and TXT formats are supported.

Responsibilities:
    - Discover supported documents recursively.
    - Validate document formats.
    - Extract source content.
    - Preserve DOCX paragraph/table order.
    - Preserve basic source information for downstream processing.

Notes:
    - JSON/JSONL metadata and evaluation files are not knowledge sources.
    - Detailed metadata enrichment is handled by metadata_manager.py.
    - DOCX paragraphs and tables are read in their real document order.
================================================================================
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Iterator, List

from docx import Document as DocxDocument
from docx.document import Document as _Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from pypdf import PdfReader

from src.schemas.document import SourceDocument


logger = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
}


class UnsupportedDocumentTypeError(ValueError):
    """Raised when an unsupported document format is encountered."""


class DocumentLoadError(RuntimeError):
    """Raised when a supported source document cannot be loaded."""


class DocumentLoader:
    """
    Discovers and loads OnlyNativeIQ authoritative source documents.

    Supported source formats:
    - PDF
    - DOCX
    - TXT

    The loader is intentionally responsible only for:
    - file discovery
    - format validation
    - basic text extraction
    - basic source metadata

    Detailed metadata enrichment belongs in metadata_manager.py.
    """

    def __init__(
        self,
        knowledge_base_path: str | Path = "knowledge_base",
    ) -> None:
        self.knowledge_base_path = Path(knowledge_base_path)

    def discover_documents(self) -> List[Path]:
        """
        Recursively discover all supported source documents.

        Files such as JSON, JSONL, ZIP, images, etc. are ignored.
        """

        if not self.knowledge_base_path.exists():
            raise FileNotFoundError(
                f"Knowledge base folder does not exist: "
                f"{self.knowledge_base_path.resolve()}"
            )

        discovered: List[Path] = []

        for file_path in self.knowledge_base_path.rglob("*"):
            if not file_path.is_file():
                continue

            if file_path.name.startswith("."):
                continue

            extension = file_path.suffix.lower()

            if extension in SUPPORTED_EXTENSIONS:
                discovered.append(file_path)

        discovered.sort()

        logger.info(
            "Discovered %s supported knowledge documents under %s",
            len(discovered),
            self.knowledge_base_path,
        )

        return discovered

    def load_all(self) -> List[SourceDocument]:
        """
        Discover and load all supported knowledge-base documents.
        """

        paths = self.discover_documents()

        documents: List[SourceDocument] = []

        for path in paths:
            try:
                document = self.load(path)
                documents.append(document)

            except Exception:
                logger.exception(
                    "Failed to load knowledge document: %s",
                    path,
                )
                raise

        return documents

    def load(self, file_path: str | Path) -> SourceDocument:
        """
        Load one source document.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"Document does not exist: {path}"
            )

        if not path.is_file():
            raise ValueError(
                f"Expected a file but received: {path}"
            )

        extension = path.suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentTypeError(
                f"Unsupported document type '{extension}' "
                f"for file: {path.name}. "
                f"Supported formats are PDF, DOCX and TXT."
            )

        try:
            if extension == ".pdf":
                content = self._load_pdf(path)

            elif extension == ".docx":
                content = self._load_docx(path)

            elif extension == ".txt":
                content = self._load_txt(path)

            else:
                raise UnsupportedDocumentTypeError(
                    f"Unsupported extension: {extension}"
                )

        except Exception as exc:
            raise DocumentLoadError(
                f"Failed to load document '{path.name}': {exc}"
            ) from exc

        return SourceDocument(
            document_id=None,
            file_name=path.name,
            file_path=str(path.resolve()),
            file_extension=extension,
            domain=self._infer_domain(path),
            content=content,
            metadata={
                "relative_path": self._get_relative_path(path),
            },
        )

    def _load_pdf(self, path: Path) -> str:
        """
        Extract plain text from a PDF.

        Page-specific structure will be enhanced later in
        document_parser.py.
        """

        reader = PdfReader(str(path))

        page_texts: List[str] = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            if text.strip():
                page_texts.append(
                    f"[PAGE {page_number}]\n{text.strip()}"
                )

        return "\n\n".join(page_texts)

    def _load_docx(self, path: Path) -> str:
        """
        Extract DOCX paragraphs and tables in their actual document order.

        Important:
            document.paragraphs and document.tables are separate collections
            and therefore do NOT preserve the interleaved order in which
            paragraphs and tables appear in the Word file.

        This implementation walks the underlying document body so that a
        structure such as:

            Product Facts
            <table>
            Standard Batch Formulation
            <table>

        remains in that same order after extraction.
        """

        document = DocxDocument(str(path))

        parts: List[str] = []

        for block in self._iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                text = block.text.strip()

                if text:
                    parts.append(text)

            elif isinstance(block, Table):
                table_text = self._table_to_text(block)

                if table_text:
                    parts.append(table_text)

        return "\n\n".join(parts)

    @staticmethod
    def _iter_docx_blocks(
        document: _Document,
    ) -> Iterator[Paragraph | Table]:
        """
        Yield Paragraph and Table objects in actual DOCX body order.

        python-docx exposes document.paragraphs and document.tables
        independently, so iterating those collections separately loses the
        original sequence. This method walks the XML body children directly.
        """

        parent_element = document.element.body

        for child in parent_element.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(
                    child,
                    document,
                )

            elif isinstance(child, CT_Tbl):
                yield Table(
                    child,
                    document,
                )

    @staticmethod
    def _table_to_text(
        table: Table,
    ) -> str:
        """
        Convert one DOCX table to pipe-separated plain text.

        Each Word table row becomes one text line, preserving row order.

        Example:
            Shelf Life | 30 days
            Storage | Dry and cool place
        """

        table_rows: List[str] = []

        for row in table.rows:
            values = [
                DocumentLoader._clean_cell_text(
                    cell.text
                )
                for cell in row.cells
            ]

            if any(values):
                table_rows.append(
                    " | ".join(values)
                )

        return "\n".join(table_rows)

    @staticmethod
    def _clean_cell_text(
        value: str,
    ) -> str:
        """
        Normalize text extracted from a Word table cell.

        Word cells may contain paragraph breaks. They are collapsed to spaces
        so one logical table row remains one output line.
        """

        lines = [
            line.strip()
            for line in value.splitlines()
            if line.strip()
        ]

        return " ".join(lines)

    def _load_txt(self, path: Path) -> str:
        """
        Load a UTF-8 TXT file.

        utf-8-sig also handles files containing a UTF-8 BOM.
        """

        return path.read_text(
            encoding="utf-8-sig"
        ).strip()

    def _infer_domain(self, path: Path) -> str:
        """
        Infer knowledge domain from the folder immediately beneath
        knowledge_base/.

        Example:

        knowledge_base/products/Classic_Thekua.docx

        becomes:

        domain = "products"
        """

        try:
            relative_path = path.relative_to(
                self.knowledge_base_path
            )

        except ValueError:
            return "unknown"

        if len(relative_path.parts) > 1:
            return relative_path.parts[0]

        return "uncategorized"

    def _get_relative_path(self, path: Path) -> str:
        try:
            return str(
                path.relative_to(
                    self.knowledge_base_path
                )
            )

        except ValueError:
            return path.name
